"""Test backend_data_extract (text/table kiểu documents-management/backend_data)."""

import io

import httpx
import pytest

from raganything.backend_data_extract import (
    BackendDataExtractor,
    BackendDataParseError,
    append_image_items,
    build_pdf_content_blocks,
    rebuild_content_list,
    rebuild_docx_content_list,
)


def _plumber_block(content, page, top, btype="paragraph"):
    return {
        "type": btype,
        "content": content,
        "page": page,
        "bbox": {"x0": 0, "top": top, "x1": 10, "bottom": top + 5},
    }


def test_parse_table_response_formats():
    ex = BackendDataExtractor()
    assert ex._parse_table_response("<table></table>")[0]["html"] == "<table></table>"
    data = {"tables": [{"html": "<table><tr><td>a</td></tr></table>", "page": 1}]}
    parsed = ex._parse_table_response(data)
    assert parsed[0]["html"].startswith("<table>")
    assert parsed[0]["page"] == 1
    parsed = ex._parse_table_response([{"html": "<t/>", "page": 2}])
    assert parsed[0]["page"] == 2


def test_build_pdf_content_blocks_replaces_tables_with_qwen_html():
    blocks = [
        _plumber_block("Tiêu đề", 1, 10),
        _plumber_block("| a | b |", 1, 20, btype="table"),
        _plumber_block("Kết luận", 1, 30),
    ]
    qwen = [{"html": "<table><tr><td>x</td></tr></table>", "page": 1}]
    out = build_pdf_content_blocks(blocks, qwen)
    assert [b["type"] for b in out] == ["text", "table", "text"]
    assert out[0]["text"] == "Tiêu đề"
    assert out[1]["table_body"].startswith("<table><tr><td>x</td>")
    assert out[1]["page_idx"] == 0
    assert out[2]["text"] == "Kết luận"


def test_build_pdf_content_blocks_falls_back_to_plumber_html():
    blocks = [_plumber_block("| a | b |", 2, 20, btype="table")]
    out = build_pdf_content_blocks(blocks, [])
    assert out[0]["type"] == "table"
    assert out[0]["table_body"] == "| a | b |"
    assert out[0]["page_idx"] == 1


def test_build_pdf_content_blocks_zero_based_qwen_pages():
    blocks = [
        _plumber_block("Đoạn 1", 1, 10),
        _plumber_block("| a | b |", 1, 20, btype="table"),
    ]
    qwen = [{"html": "<table><tr><td>q</td></tr></table>", "page": 0}]
    out = build_pdf_content_blocks(blocks, qwen)
    assert out[1]["table_body"].startswith("<table><tr><td>q</td>")


def test_build_pdf_content_blocks_leftover_qwen_tables():
    blocks = [_plumber_block("Chỉ text", 1, 10)]
    qwen = [
        {"html": "<table><tr><td>1</td></tr></table>", "page": 1},
        {"html": "<table><tr><td>2</td></tr></table>", "page": 3},
        {"html": "<table><tr><td>3</td></tr></table>", "page": None},
    ]
    out = build_pdf_content_blocks(blocks, qwen)
    tables = [b for b in out if b["type"] == "table"]
    assert len(tables) == 3
    assert tables[0]["page_idx"] == 0
    assert tables[1]["page_idx"] == 2
    assert tables[2]["page_idx"] is None


def test_append_image_items_pdf_same_page():
    blocks = [
        {"type": "text", "text": "a", "page_idx": 0},
        {"type": "text", "text": "b", "page_idx": 1},
    ]
    images = [
        {"type": "image", "img_path": "p1.png", "page_idx": 0},
        {"type": "image", "img_path": "p2.png", "page_idx": 1},
    ]
    out = append_image_items(blocks, images)
    assert [(b["type"], b.get("page_idx")) for b in out] == [
        ("text", 0),
        ("image", 0),
        ("text", 1),
        ("image", 1),
    ]


def test_append_image_items_docx_pages_after_text():
    blocks = [
        {"type": "text", "text": "mở đầu", "page_idx": None},
        {"type": "table", "table_body": "<table/>", "page_idx": None},
    ]
    images = [
        {"type": "image", "img_path": "a.png", "page_idx": 2},
        {"type": "image", "img_path": "b.png", "page_idx": 1},
    ]
    out = append_image_items(blocks, images)
    assert [(b["type"], b.get("page_idx")) for b in out] == [
        ("text", None),
        ("table", None),
        ("image", 1),
        ("image", 2),
    ]


def test_extract_docx_blocks_mapping():
    from docx import Document

    doc = Document()
    doc.add_heading("Giới thiệu", level=1)
    doc.add_paragraph("Nội dung tiếng Việt")
    from docx.oxml.ns import qn

    p = doc.add_paragraph("Mục con")
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("w:numPr"), {}))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tên"
    table.cell(0, 1).text = "Tuổi"
    table.cell(1, 0).text = "An"
    table.cell(1, 1).text = "20"

    buf = io.BytesIO()
    doc.save(buf)
    ex = BackendDataExtractor()
    blocks = ex.extract_docx_blocks(buf.getvalue())

    assert [b["type"] for b in blocks] == ["heading", "paragraph", "list", "table"]
    assert blocks[0]["content"] == "# Giới thiệu"
    assert blocks[1]["content"] == "Nội dung tiếng Việt"
    assert blocks[2]["content"] == "- Mục con"
    assert blocks[3]["html"].count("<th>") == 2
    assert blocks[3]["html"].count("<td>") == 2


async def test_rebuild_docx_content_list_merges_mineru_images():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Đoạn văn")
    buf = io.BytesIO()
    doc.save(buf)

    mineru_list = [
        {"type": "text", "text": "mineru text", "page_idx": 0},
        {"type": "table", "table_body": "<table/>", "page_idx": 0},
        {"type": "equation", "text": "E=mc2", "page_idx": 1},
        {"type": "image", "img_path": "img.png", "page_idx": 1},
    ]
    ex = BackendDataExtractor()
    out = await rebuild_docx_content_list(mineru_list, buf.getvalue(), ex)
    assert [b["type"] for b in out] == ["text", "image"]
    assert out[0]["text"] == "Đoạn văn"
    assert out[0]["page_idx"] is None
    assert out[1]["img_path"] == "img.png"


async def test_rebuild_content_list_dispatch(tmp_path):
    ex = BackendDataExtractor()
    txt = tmp_path / "a.txt"
    txt.write_text("hello", encoding="utf-8")
    assert await rebuild_content_list([], txt, ex) is None

    docx = tmp_path / "b.docx"
    from docx import Document

    doc = Document()
    doc.add_paragraph("Xin chào")
    doc.save(docx)
    out = await rebuild_content_list([], docx, ex)
    assert out and out[0]["text"] == "Xin chào"


async def test_rebuild_pdf_content_list_full_pipeline(tmp_path, monkeypatch):
    import fitz

    pdf_path = tmp_path / "doc.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Trang mot tieng Viet")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "Trang hai noi dung")
    doc.save(pdf_path)
    doc.close()

    def fake_post(self, endpoint, file_bytes, filename, content_type):
        if "qwen" in endpoint:
            return {"tables": [{"html": "<table><tr><td>q</td></tr></table>", "page": 1}]}
        return {"text": "", "content": {"text": "ocr fallback", "children": []}}

    monkeypatch.setattr(BackendDataExtractor, "_post_file", fake_post)
    ex = BackendDataExtractor()

    mineru_list = [
        {"type": "text", "text": "mineru text", "page_idx": 0},
        {"type": "image", "img_path": "img.png", "page_idx": 1},
    ]
    out = await rebuild_content_list(mineru_list, pdf_path, ex)

    types = [b["type"] for b in out]
    assert types.count("text") == 2
    assert types.count("image") == 1
    assert "mineru text" not in [b.get("text", "") for b in out if b["type"] == "text"]
    texts = [b["text"] for b in out if b["type"] == "text"]
    assert "Trang mot tieng Viet" in texts[0]
    assert "Trang hai noi dung" in texts[1]
    assert out[-1]["type"] == "image" and out[-1]["page_idx"] == 1


async def test_extract_tables_via_mock_transport():
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path.endswith("/api/v1/raw_miner_ocr/qwen")
        return httpx.Response(200, json={"tables": [{"html": "<table/>", "page": 1}]})

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    ex = BackendDataExtractor(base_url="http://ocr-fake:12006", http_client=client)
    tables = await ex.extract_tables(b"%PDF-fake", "doc.pdf")
    assert tables[0]["page"] == 1
    await client.aclose()


async def test_extract_scan_text_error_raises():
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(500, text="boom")

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    ex = BackendDataExtractor(base_url="http://ocr-fake:12006", http_client=client)
    with pytest.raises(BackendDataParseError):
        await ex.extract_scan_text(b"%PDF-fake", "doc.pdf")
    await client.aclose()
